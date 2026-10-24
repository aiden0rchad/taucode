import importlib.util
import os
from pathlib import Path
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("render", ROOT / "skills/render.py")
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class DocumentTests(unittest.TestCase):
    def test_skill_frontmatter(self):
        for name in ("docx", "xlsx", "pptx", "pdf", "charts"):
            text = (ROOT / "skills" / name / "SKILL.md").read_text()
            front = text.split("---", 2)[1]
            self.assertIn(f"name: {name}", front)
            self.assertIn("description: ", front)
            self.assertIn("unavailable", text)

    def test_boundaries(self):
        with self.assertRaises(ValueError):
            module.render("docx", {}, "../escape.docx")

    @unittest.skipUnless(all(importlib.util.find_spec(m) for m in ["docx"]), "Document libraries unavailable")
    def test_docx(self):
        previous = Path.cwd()
        with tempfile.TemporaryDirectory() as temp:
            try:
                os.chdir(temp)
                path = module.render("docx", {"title":"Brief","paragraphs":["Summary"]}, "outputs/result.docx")
                self.assertGreater(path.stat().st_size, 100)
                from docx import Document
                self.assertEqual(Document(path).paragraphs[0].text, "Brief")
            finally:
                os.chdir(previous)

    @unittest.skipUnless(all(importlib.util.find_spec(m) for m in ["openpyxl"]), "Document libraries unavailable")
    def test_xlsx(self):
        previous = Path.cwd()
        with tempfile.TemporaryDirectory() as temp:
            try:
                os.chdir(temp)
                path = module.render("xlsx", {"rows":[["Month","Total"],["January",42]]}, "outputs/result.xlsx")
                self.assertGreater(path.stat().st_size, 100)
                from openpyxl import load_workbook
                self.assertEqual(load_workbook(path).active["B2"].value, 42)
            finally:
                os.chdir(previous)

    @unittest.skipUnless(all(importlib.util.find_spec(m) for m in ["pptx"]), "Document libraries unavailable")
    def test_pptx(self):
        previous = Path.cwd()
        with tempfile.TemporaryDirectory() as temp:
            try:
                os.chdir(temp)
                path = module.render("pptx", {"slides":[{"title":"Summary","bullets":["First point"]}]}, "outputs/result.pptx")
                self.assertGreater(path.stat().st_size, 100)
                from pptx import Presentation
                self.assertEqual(len(Presentation(path).slides), 1)
            finally:
                os.chdir(previous)

    @unittest.skipUnless(all(importlib.util.find_spec(m) for m in ["pypdf","reportlab"]), "Document libraries unavailable")
    def test_pdf(self):
        previous = Path.cwd()
        with tempfile.TemporaryDirectory() as temp:
            try:
                os.chdir(temp)
                path = module.render("pdf", {"title":"Brief","paragraphs":["Summary"]}, "outputs/result.pdf")
                self.assertGreater(path.stat().st_size, 100)
                from pypdf import PdfReader
                self.assertIn("Summary", PdfReader(path).pages[0].extract_text())
            finally:
                os.chdir(previous)

    @unittest.skipUnless(all(importlib.util.find_spec(m) for m in ["matplotlib"]), "Document libraries unavailable")
    def test_charts(self):
        previous = Path.cwd()
        with tempfile.TemporaryDirectory() as temp:
            try:
                os.chdir(temp)
                path = module.render("charts", {"title":"Totals","labels":["January"],"values":[42],"ylabel":"USD"}, "outputs/result.png")
                self.assertGreater(path.stat().st_size, 100)
                self.assertEqual(path.read_bytes()[:8], b"\x89PNG\r\n\x1a\n")
            finally:
                os.chdir(previous)


