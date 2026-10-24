"""Document starters; output paths stay inside the current space's outputs/."""
import argparse
import json
from pathlib import Path


def render(kind, data, destination):
    path = Path(destination).resolve()
    root = (Path.cwd() / "outputs").resolve()
    if not root.is_relative_to(Path.cwd().resolve()) or not path.is_relative_to(root):
        raise ValueError("Destination must be inside outputs/")
    path.parent.mkdir(parents=True, exist_ok=True)
    if kind == "docx":
        from docx import Document
        document = Document()
        document.add_heading(data["title"], 0)
        for paragraph in data["paragraphs"]:
            document.add_paragraph(paragraph)
        document.save(path)
    elif kind == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font
        book = Workbook()
        sheet = book.active
        sheet.title = "Summary"
        for row in data["rows"]:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            sheet.column_dimensions[cell.column_letter].width = 24
        sheet.freeze_panes = "A2"
        book.save(path)
    elif kind == "pptx":
        from pptx import Presentation
        deck = Presentation()
        for item in data["slides"]:
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = item["title"]
            slide.placeholders[1].text = "\n".join(item["bullets"])
        deck.save(path)
    elif kind == "pdf":
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from xml.sax.saxutils import escape
        styles = getSampleStyleSheet()
        story = [Paragraph(escape(data["title"]), styles["Title"])]
        for text in data["paragraphs"]:
            story.extend([Paragraph(escape(text), styles["BodyText"]), Spacer(1, 12)])
        SimpleDocTemplate(str(path)).build(story)
    elif kind == "charts":
        import matplotlib
        matplotlib.use("Agg")
        from matplotlib import pyplot as plt
        figure, axis = plt.subplots(layout="constrained")
        axis.bar(data["labels"], data["values"])
        axis.set(title=data["title"], xlabel=data.get("xlabel", ""), ylabel=data.get("ylabel", ""))
        figure.savefig(path, dpi=180)
        plt.close(figure)
    else:
        raise ValueError("Unknown document kind")
    return path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("kind", choices=["docx", "xlsx", "pptx", "pdf", "charts"])
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    try:
        render(args.kind, json.loads(args.input.read_text(encoding="utf-8")), args.output)
    except ImportError as error:
        parser.exit(2, f"Document library unavailable: {error}. Use Markdown/CSV or retry Work runtime setup.\n")

